"""
Output file generation (sections 31-33). Streams rows out of the staging
table in chunks -- never builds an in-memory list of the full result set,
and never uses pandas for this. CSV via the stdlib `csv` module; the
consolidated Excel workbook via `openpyxl.Workbook(write_only=True)`,
automatically starting a new sheet if a category would exceed Excel's
1,048,576-row-per-sheet limit.
"""
from __future__ import annotations

import csv
import logging
from pathlib import Path

import openpyxl
import pyodbc

from app.filtration.reason_codes import KEPT_CODE, display_name, output_filename
from app.models.job import JobColumn
from app.utils.identifiers import quote_alias, quote_ident

logger = logging.getLogger("ft.report_service")

EXCEL_MAX_ROWS_PER_SHEET = 1_048_575  # leave room for the header row
FETCH_CHUNK_SIZE = 20_000


def _select_columns_sql(job_columns: list[JobColumn]) -> tuple[str, list[str]]:
    """Builds `s.[c_0_email] AS [Email], ...` plus the display header list,
    so output files show the ORIGINAL uploaded header names, not the
    sanitized internal SQL column names."""
    select_parts = []
    headers = []
    for jc in job_columns:
        select_parts.append(f"s.{quote_ident(jc.SqlColumnName)} AS {quote_alias(jc.OriginalName)}")
        headers.append(jc.OriginalName)
    return ", ".join(select_parts), headers


def _where_clause_for_reason(reason_code: str) -> str:
    if reason_code == KEPT_CODE:
        return "s.ReasonCode IS NULL AND s.IsOtherTLD = 0"
    return "s.ReasonCode = ?"


def _stream_category_rows(cursor: pyodbc.Cursor, staging_table: str, select_sql: str, reason_code: str):
    where_sql = _where_clause_for_reason(reason_code)
    sql = f"SELECT {select_sql} FROM {quote_ident(staging_table)} AS s WHERE {where_sql}"
    if reason_code == KEPT_CODE:
        cursor.execute(sql)
    else:
        cursor.execute(sql, [reason_code])
    while True:
        rows = cursor.fetchmany(FETCH_CHUNK_SIZE)
        if not rows:
            break
        for row in rows:
            yield list(row)


def generate_category_csv(
    conn: pyodbc.Connection,
    staging_table: str,
    job_columns: list[JobColumn],
    reason_code: str,
    output_dir: Path,
) -> Path | None:
    select_sql, headers = _select_columns_sql(job_columns)
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / output_filename(reason_code)

    cursor = conn.cursor()
    row_iter = _stream_category_rows(cursor, staging_table, select_sql, reason_code)

    wrote_any = False
    with open(out_path, "w", newline="", encoding="utf-8-sig") as f:
        writer = csv.writer(f)
        writer.writerow(headers)
        for row in row_iter:
            writer.writerow(row)
            wrote_any = True

    if not wrote_any:
        out_path.unlink(missing_ok=True)
        return None
    return out_path


def generate_all_category_csvs(
    conn: pyodbc.Connection,
    staging_table: str,
    job_columns: list[JobColumn],
    summary: dict[str, int],
    output_dir: Path,
) -> dict[str, Path]:
    """Only creates files for categories that actually contain records
    (section 31)."""
    results: dict[str, Path] = {}
    for reason_code, count in summary.items():
        if count <= 0:
            continue
        path = generate_category_csv(conn, staging_table, job_columns, reason_code, output_dir)
        if path:
            results[reason_code] = path
    return results


def generate_consolidated_excel(
    conn: pyodbc.Connection,
    staging_table: str,
    job_columns: list[JobColumn],
    summary: dict[str, int],
    output_dir: Path,
    job_meta: dict,
) -> Path:
    select_sql, headers = _select_columns_sql(job_columns)
    output_dir.mkdir(parents=True, exist_ok=True)
    out_path = output_dir / "Report.xlsx"

    wb = openpyxl.Workbook(write_only=True)

    summary_ws = wb.create_sheet("Summary")
    summary_ws.append(["Job ID", job_meta.get("job_id")])
    summary_ws.append(["File Name", job_meta.get("file_name")])
    summary_ws.append(["Processed At", job_meta.get("processed_at")])
    summary_ws.append([])
    summary_ws.append(["Category", "Row Count"])
    summary_ws.append(["INPUT", job_meta.get("total_rows", 0)])
    for reason_code, count in summary.items():
        summary_ws.append([display_name(reason_code), count])

    cursor = conn.cursor()
    for reason_code, count in summary.items():
        if count <= 0:
            continue
        sheet_title = display_name(reason_code)[:31]
        ws = wb.create_sheet(sheet_title)
        ws.append(headers)
        sheet_part = 1
        rows_in_sheet = 0
        for row in _stream_category_rows(cursor, staging_table, select_sql, reason_code):
            if rows_in_sheet >= EXCEL_MAX_ROWS_PER_SHEET:
                sheet_part += 1
                ws = wb.create_sheet(f"{sheet_title[:28]}_{sheet_part}")
                ws.append(headers)
                rows_in_sheet = 0
            ws.append(row)
            rows_in_sheet += 1

    wb.save(str(out_path))
    return out_path
