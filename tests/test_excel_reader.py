"""Verifies every worksheet in a multi-sheet workbook is read (section 62:
'Create a test workbook with at least Sheet1, Sheet2, Sheet3 and verify
every sheet is processed'). Uses openpyxl to build a real .xlsx fixture on
disk rather than mocking anything."""
import openpyxl
import pytest

from app.filtration.excel_reader import iter_all_sheet_rows, sheet_names, unified_headers


@pytest.fixture
def multi_sheet_workbook(tmp_path):
    path = tmp_path / "multi_sheet.xlsx"
    wb = openpyxl.Workbook()

    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["Email", "Name", "Title"])
    ws1.append(["john@company.com", "John Smith", "Manager"])
    ws1.append(["jane@company.com", "Jane Doe", "Director"])

    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["Email", "Name", "Title"])
    ws2.append(["mike@company.com", "Mike Ross", "Analyst"])

    ws3 = wb.create_sheet("Sheet3")
    ws3.append(["Email", "Name", "Title"])
    ws3.append(["sara@company.com", "Sara Lee", "VP"])
    ws3.append(["tom@company.com", "Tom Hardy", "Engineer"])

    wb.save(path)
    return path


def test_all_sheet_names_present(multi_sheet_workbook):
    assert sheet_names(multi_sheet_workbook) == ["Sheet1", "Sheet2", "Sheet3"]


def test_unified_headers_from_all_sheets(multi_sheet_workbook):
    headers = unified_headers(multi_sheet_workbook)
    assert headers == ["Email", "Name", "Title"]


def test_every_sheet_is_processed_not_just_sheet1(multi_sheet_workbook):
    headers = unified_headers(multi_sheet_workbook)
    rows = list(iter_all_sheet_rows(multi_sheet_workbook, headers))

    sheets_seen = {sheet for sheet, _, _ in rows}
    assert sheets_seen == {"Sheet1", "Sheet2", "Sheet3"}

    total_data_rows = len(rows)
    assert total_data_rows == 5  # 2 + 1 + 2

    emails = {values[0] for _, _, values in rows}
    assert emails == {
        "john@company.com",
        "jane@company.com",
        "mike@company.com",
        "sara@company.com",
        "tom@company.com",
    }


def test_rows_aligned_to_unified_header_order(multi_sheet_workbook):
    headers = unified_headers(multi_sheet_workbook)
    rows = list(iter_all_sheet_rows(multi_sheet_workbook, headers))
    sheet1_row = next(v for s, _, v in rows if s == "Sheet1" and v[0] == "john@company.com")
    assert sheet1_row == ["john@company.com", "John Smith", "Manager"]


def test_workbook_with_mismatched_columns_across_sheets(tmp_path):
    path = tmp_path / "mismatched.xlsx"
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1.append(["Email", "Name"])
    ws1.append(["a@x.com", "A"])

    ws2 = wb.create_sheet("Sheet2")
    ws2.append(["Email", "Company"])  # different second column
    ws2.append(["b@x.com", "Acme"])
    wb.save(path)

    headers = unified_headers(path)
    assert set(headers) == {"Email", "Name", "Company"}

    rows = list(iter_all_sheet_rows(path, headers))
    assert len(rows) == 2
    row_a = next(v for s, _, v in rows if s == "Sheet1")
    row_b = next(v for s, _, v in rows if s == "Sheet2")
    name_idx = headers.index("Name")
    company_idx = headers.index("Company")
    assert row_a[name_idx] == "A" and row_a[company_idx] is None
    assert row_b[company_idx] == "Acme" and row_b[name_idx] is None
